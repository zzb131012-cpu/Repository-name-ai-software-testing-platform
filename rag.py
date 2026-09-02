import os
import json
import hashlib
import shutil

from pypdf import PdfReader
from docx import Document as DocxDocument

from langchain_chroma import Chroma
from langchain_huggingface import HuggingFaceEmbeddings
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter


# =========================================================
# 1. 路径配置
# =========================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

KNOWLEDGE_DIR = os.path.join(
    BASE_DIR,
    "knowledge"
)

CHROMA_DIR = os.path.join(
    BASE_DIR,
    "chroma_db"
)

HASH_FILE = os.path.join(
    BASE_DIR,
    "knowledge_hash.json"
)

os.makedirs(
    KNOWLEDGE_DIR,
    exist_ok=True
)


# =========================================================
# 2. Embedding 模型
# =========================================================

embeddings = HuggingFaceEmbeddings(
    model_name="BAAI/bge-small-zh-v1.5",
    model_kwargs={
        "device": "cpu"
    },
    encode_kwargs={
        "normalize_embeddings": True
    }
)


# =========================================================
# 3. 获取知识库文件
# =========================================================

def get_knowledge_files():

    files = []

    supported_extensions = (
        ".txt",
        ".md",
        ".pdf",
        ".docx"
    )

    for file_name in sorted(
        os.listdir(KNOWLEDGE_DIR)
    ):

        file_path = os.path.join(
            KNOWLEDGE_DIR,
            file_name
        )

        if not os.path.isfile(
            file_path
        ):
            continue

        if file_name.lower().endswith(
            supported_extensions
        ):
            files.append(
                file_path
            )

    return files


# =========================================================
# 4. 获取知识库文件名列表
# =========================================================

def get_knowledge_file_list():

    files = get_knowledge_files()

    return [
        os.path.basename(file_path)
        for file_path in files
    ]


# =========================================================
# 5. TXT / Markdown 读取
# =========================================================

def read_text_file(
    file_path
):

    try:

        with open(
            file_path,
            "r",
            encoding="utf-8"
        ) as f:

            return f.read()

    except UnicodeDecodeError:

        with open(
            file_path,
            "r",
            encoding="gbk",
            errors="ignore"
        ) as f:

            return f.read()


# =========================================================
# 6. PDF 读取
# =========================================================

def read_pdf_file(
    file_path
):

    text = ""

    reader = PdfReader(
        file_path
    )

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):

        page_text = (
            page.extract_text()
        )

        if page_text:

            text += (
                f"\n===== 第 {page_number} 页 =====\n"
            )

            text += page_text

    return text


# =========================================================
# 7. Word 读取
# =========================================================

def read_docx_file(
    file_path
):

    text = ""

    document = DocxDocument(
        file_path
    )

    for paragraph in (
        document.paragraphs
    ):

        paragraph_text = (
            paragraph.text.strip()
        )

        if paragraph_text:

            text += (
                paragraph_text
                + "\n"
            )

    for table in (
        document.tables
    ):

        text += (
            "\n===== 表格 =====\n"
        )

        for row in table.rows:

            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            text += (
                " | ".join(values)
                + "\n"
            )

    return text


# =========================================================
# 8. 统一读取知识库文件
# =========================================================

def read_knowledge_file(
    file_path
):

    lower_name = (
        file_path.lower()
    )

    if lower_name.endswith(
        (".txt", ".md")
    ):

        return read_text_file(
            file_path
        )

    if lower_name.endswith(
        ".pdf"
    ):

        return read_pdf_file(
            file_path
        )

    if lower_name.endswith(
        ".docx"
    ):

        return read_docx_file(
            file_path
        )

    return ""


# =========================================================
# 9. 根据文件名查看内容
# =========================================================

def get_knowledge_file_content(
    file_name
):

    safe_name = (
        os.path.basename(
            file_name
        )
    )

    file_path = os.path.join(
        KNOWLEDGE_DIR,
        safe_name
    )

    if not os.path.exists(
        file_path
    ):

        return ""

    try:

        return read_knowledge_file(
            file_path
        )

    except Exception as e:

        return (
            f"文件读取失败：{e}"
        )


# =========================================================
# 10. 删除知识库文件
# =========================================================

def delete_knowledge_file(
    file_name
):

    safe_name = (
        os.path.basename(
            file_name
        )
    )

    file_path = os.path.join(
        KNOWLEDGE_DIR,
        safe_name
    )

    if not os.path.exists(
        file_path
    ):

        return False

    try:

        os.remove(
            file_path
        )

        return True

    except Exception as e:

        print(
            f"删除知识库文件失败：{e}"
        )

        return False


# =========================================================
# 11. 计算知识库 Hash
# =========================================================

def calculate_knowledge_hash():

    hasher = hashlib.sha256()

    files = (
        get_knowledge_files()
    )

    for file_path in files:

        file_name = (
            os.path.basename(
                file_path
            )
        )

        hasher.update(
            file_name.encode(
                "utf-8"
            )
        )

        with open(
            file_path,
            "rb"
        ) as f:

            while True:

                chunk = f.read(
                    8192
                )

                if not chunk:
                    break

                hasher.update(
                    chunk
                )

    return hasher.hexdigest()


# =========================================================
# 12. 读取旧 Hash
# =========================================================

def load_saved_hash():

    if not os.path.exists(
        HASH_FILE
    ):

        return None

    try:

        with open(
            HASH_FILE,
            "r",
            encoding="utf-8"
        ) as f:

            data = json.load(
                f
            )

        return data.get(
            "knowledge_hash"
        )

    except Exception:

        return None


# =========================================================
# 13. 保存 Hash
# =========================================================

def save_hash(
    current_hash
):

    with open(
        HASH_FILE,
        "w",
        encoding="utf-8"
    ) as f:

        json.dump(
            {
                "knowledge_hash":
                    current_hash
            },
            f,
            ensure_ascii=False,
            indent=2
        )


# =========================================================
# 14. 转换为 LangChain Document
# =========================================================

def load_knowledge_documents():

    documents = []

    for file_path in (
        get_knowledge_files()
    ):

        try:

            content = (
                read_knowledge_file(
                    file_path
                )
            )

        except Exception as e:

            print(
                f"读取失败：{file_path}"
            )

            print(
                e
            )

            continue

        if not content.strip():

            continue

        file_name = (
            os.path.basename(
                file_path
            )
        )

        document = Document(
            page_content=content,
            metadata={
                "source": file_name
            }
        )

        documents.append(
            document
        )

    return documents


# =========================================================
# 15. 文档切分
# =========================================================

def split_documents(
    documents
):

    splitter = (
        RecursiveCharacterTextSplitter(
            chunk_size=500,
            chunk_overlap=100
        )
    )

    chunks = (
        splitter.split_documents(
            documents
        )
    )

    return chunks


# =========================================================
# 16. 删除旧向量库
# =========================================================

def delete_old_vector_db():

    if os.path.exists(
        CHROMA_DIR
    ):

        shutil.rmtree(
            CHROMA_DIR
        )


# =========================================================
# 17. 创建向量数据库
# =========================================================

def build_vector_db():

    documents = (
        load_knowledge_documents()
    )

    if not documents:

        delete_old_vector_db()

        current_hash = (
            calculate_knowledge_hash()
        )

        save_hash(
            current_hash
        )

        print(
            "当前知识库为空。"
        )

        return None

    chunks = (
        split_documents(
            documents
        )
    )

    print(
        f"知识库文件：{len(documents)} 个"
    )

    print(
        f"文本块：{len(chunks)} 个"
    )

    print(
        "正在构建 RAG 知识库..."
    )

    delete_old_vector_db()

    vector_db = (
        Chroma.from_documents(
            documents=chunks,
            embedding=embeddings,
            persist_directory=CHROMA_DIR,
            collection_name="testing_knowledge"
        )
    )

    current_hash = (
        calculate_knowledge_hash()
    )

    save_hash(
        current_hash
    )

    print(
        "RAG 知识库构建完成。"
    )

    return vector_db


# =========================================================
# 18. 强制重建
# =========================================================

def rebuild_vector_db():

    return build_vector_db()


# =========================================================
# 19. 加载已有 Chroma
# =========================================================

def load_vector_db():

    if not os.path.exists(
        CHROMA_DIR
    ):

        return None

    try:

        vector_db = Chroma(
            persist_directory=CHROMA_DIR,
            embedding_function=embeddings,
            collection_name="testing_knowledge"
        )

        return vector_db

    except Exception as e:

        print(
            f"加载 Chroma 失败：{e}"
        )

        return None


# =========================================================
# 20. 自动检查知识库
# =========================================================

def ensure_vector_db():

    current_hash = (
        calculate_knowledge_hash()
    )

    saved_hash = (
        load_saved_hash()
    )

    if current_hash != saved_hash:

        print(
            "检测到知识库发生变化，自动重建。"
        )

        return build_vector_db()

    if not os.path.exists(
        CHROMA_DIR
    ):

        if get_knowledge_files():

            print(
                "没有发现 Chroma，自动创建。"
            )

            return build_vector_db()

        return None

    return load_vector_db()


# =========================================================
# 21. 搜索知识库
# =========================================================

def search_knowledge(
    query,
    k=4
):

    vector_db = (
        ensure_vector_db()
    )

    if vector_db is None:

        return []

    try:

        results = (
            vector_db.similarity_search(
                query,
                k=k
            )
        )

        return results

    except Exception as e:

        print(
            f"RAG 检索失败：{e}"
        )

        return []


# =========================================================
# 22. 生成 RAG 上下文
# =========================================================

def get_rag_context(
    query,
    k=4
):

    documents = (
        search_knowledge(
            query,
            k=k
        )
    )

    if not documents:

        return (
            "当前知识库没有检索到相关资料。"
        )

    contexts = []

    for index, doc in enumerate(
        documents,
        start=1
    ):

        source = (
            doc.metadata.get(
                "source",
                "未知来源"
            )
        )

        context = f"""
===== 参考资料 {index} =====

来源：{source}

{doc.page_content}
"""

        contexts.append(
            context
        )

    return "\n".join(
        contexts
    )


# =========================================================
# 23. 单独运行测试
# =========================================================

if __name__ == "__main__":

    ensure_vector_db()

    print(
        "\n========== RAG 测试 ==========\n"
    )

    result = get_rag_context(
        "登录模块有哪些测试风险？",
        k=4
    )

    print(
        result
    )