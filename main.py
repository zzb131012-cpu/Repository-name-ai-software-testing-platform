import os
import json
from datetime import datetime

from dotenv import load_dotenv
from langchain_openai import ChatOpenAI
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side


# =========================
# 1. 读取 .env
# =========================

load_dotenv()

api_key = os.getenv("DEEPSEEK_API_KEY")

if not api_key:
    raise ValueError(
        "没有读取到 DEEPSEEK_API_KEY，请检查 .env 文件"
    )


# =========================
# 2. 创建 DeepSeek 模型
# =========================

model = ChatOpenAI(
    model="deepseek-chat",
    api_key=api_key,
    base_url="https://api.deepseek.com"
)


# =========================
# 3. 输入多行需求
# =========================

print("=" * 50)
print("AI 测试用例生成器")
print("=" * 50)

print("\n请输入需求内容：")
print("输入完成后，单独输入 END 并回车结束：\n")

lines = []

while True:
    line = input()

    if line.strip().upper() == "END":
        break

    lines.append(line)

requirement = "\n".join(lines)

if not requirement.strip():
    raise ValueError("需求内容不能为空")


# =========================
# 4. Prompt
# =========================

prompt = f"""
你是一名专业的软件测试工程师。

请严格根据下面提供的需求生成完整的软件测试用例。

要求：

1. 覆盖正常场景
2. 覆盖异常场景
3. 覆盖边界值场景
4. 覆盖核心业务规则
5. 不允许编造需求中不存在的业务规则
6. 如果需求不明确，请在相关字段标记【需求待确认】

每条测试用例必须包含：

case_id
module
title
precondition
steps
test_data
expected_result
priority

priority 只允许使用：

P0
P1
P2

只允许输出 JSON 数组。

不要输出 Markdown。
不要输出解释文字。
不要输出 ```json。

返回格式示例：

[
  {{
    "case_id": "TC001",
    "module": "登录",
    "title": "正确手机号和密码登录",
    "precondition": "用户已经注册且账号正常",
    "steps": "1.进入登录页；2.输入正确手机号；3.输入正确密码；4.点击登录",
    "test_data": "正确手机号、正确密码",
    "expected_result": "登录成功并进入首页",
    "priority": "P0"
  }}
]

需求内容：

{requirement}
"""


# =========================
# 5. 调用 DeepSeek
# =========================

print("\n正在生成测试用例，请稍候...\n")

try:
    result = model.invoke(prompt)
    content = result.content

except Exception as e:
    print("调用 DeepSeek 失败：")
    print(e)
    exit()


# =========================
# 6. 解析 JSON
# =========================

try:
    test_cases = json.loads(content)

except json.JSONDecodeError:
    print("\nAI 返回的内容不是合法 JSON。")
    print("\nAI 原始返回内容：\n")
    print(content)
    exit()


if not isinstance(test_cases, list):
    print("AI 返回的数据不是 JSON 数组。")
    exit()


if len(test_cases) == 0:
    print("AI 没有生成测试用例。")
    exit()


print(f"成功生成测试用例：{len(test_cases)} 条")


# =========================
# 7. 创建 Excel
# =========================

wb = Workbook()

ws = wb.active

ws.title = "测试用例"


# =========================
# 8. Excel 表头
# =========================

headers = [
    "用例编号",
    "测试模块",
    "测试标题",
    "前置条件",
    "操作步骤",
    "测试数据",
    "预期结果",
    "优先级"
]

ws.append(headers)


# =========================
# 9. 写入测试用例
# =========================

for case in test_cases:
    ws.append([
        case.get("case_id", ""),
        case.get("module", ""),
        case.get("title", ""),
        case.get("precondition", ""),
        case.get("steps", ""),
        case.get("test_data", ""),
        case.get("expected_result", ""),
        case.get("priority", "")
    ])


# =========================
# 10. 设置边框
# =========================

thin = Side(style="thin")

border = Border(
    left=thin,
    right=thin,
    top=thin,
    bottom=thin
)


# =========================
# 11. 表头样式
# =========================

for cell in ws[1]:
    cell.font = Font(
        bold=True,
        size=11
    )

    cell.alignment = Alignment(
        horizontal="center",
        vertical="center"
    )

    cell.border = border


# =========================
# 12. 内容样式
# =========================

for row in ws.iter_rows(min_row=2):
    for cell in row:

        cell.alignment = Alignment(
            vertical="top",
            wrap_text=True
        )

        cell.border = border


# =========================
# 13. 用例编号、模块、优先级居中
# =========================

for row in range(2, ws.max_row + 1):

    ws[f"A{row}"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    ws[f"B{row}"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )

    ws[f"H{row}"].alignment = Alignment(
        horizontal="center",
        vertical="center",
        wrap_text=True
    )


# =========================
# 14. 设置列宽
# =========================

column_widths = {
    "A": 12,
    "B": 15,
    "C": 30,
    "D": 30,
    "E": 50,
    "F": 35,
    "G": 45,
    "H": 10
}

for column, width in column_widths.items():
    ws.column_dimensions[column].width = width


# =========================
# 15. 设置行高
# =========================

ws.row_dimensions[1].height = 28

for row in range(2, ws.max_row + 1):
    ws.row_dimensions[row].height = 70


# =========================
# 16. 冻结首行
# =========================

ws.freeze_panes = "A2"


# =========================
# 17. 开启自动筛选
# =========================

ws.auto_filter.ref = ws.dimensions


# =========================
# 18. 第12步核心：
# 自动生成时间戳文件名
# =========================

current_time = datetime.now().strftime(
    "%Y%m%d_%H%M%S"
)

file_name = f"test_cases_{current_time}.xlsx"


# =========================
# 19. 保存 Excel
# =========================

try:
    wb.save(file_name)

except PermissionError:
    print("\nExcel 文件保存失败。")
    print("请检查 Excel 文件是否正在被打开。")
    exit()


# =========================
# 20. 输出结果
# =========================

print("\n" + "=" * 50)

print("测试用例生成成功！")

print(f"测试用例数量：{len(test_cases)}")

print(f"Excel 文件：{file_name}")

print("=" * 50)