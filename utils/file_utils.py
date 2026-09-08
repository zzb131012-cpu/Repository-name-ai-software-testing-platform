from pypdf import PdfReader
from docx import Document


SUPPORTED_EXTENSIONS = (
    ".txt",
    ".md",
    ".pdf",
    ".docx",
)


def read_text_file(uploaded_file):
    """
    读取 TXT / Markdown。
    """

    content = uploaded_file.getvalue()

    try:
        return content.decode("utf-8")

    except UnicodeDecodeError:
        return content.decode(
            "gbk",
            errors="ignore"
        )


def read_pdf_file(uploaded_file):
    """
    读取 PDF。
    """

    reader = PdfReader(uploaded_file)

    text_parts = []

    for page_number, page in enumerate(
        reader.pages,
        start=1
    ):
        page_text = page.extract_text()

        if not page_text:
            continue

        text_parts.append(
            f"===== 第 {page_number} 页 ====="
        )

        text_parts.append(page_text)

    return "\n".join(text_parts)


def read_word_file(uploaded_file):
    """
    读取 DOCX。
    """

    document = Document(uploaded_file)

    text_parts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            text_parts.append(text)

    for table_index, table in enumerate(
        document.tables,
        start=1
    ):
        text_parts.append(
            f"===== 表格 {table_index} ====="
        )

        for row in table.rows:
            row_text = " | ".join(
                cell.text.strip()
                for cell in row.cells
            )

            text_parts.append(row_text)

    return "\n".join(text_parts)


def get_file_extension(filename):
    """
    获取支持的扩展名。
    """

    if not filename:
        return ""

    filename = filename.lower()

    for extension in SUPPORTED_EXTENSIONS:
        if filename.endswith(extension):
            return extension

    return ""


def validate_uploaded_file(uploaded_file):
    """
    校验上传文件。

    返回：
    (是否有效, 提示信息)
    """

    if uploaded_file is None:
        return False, "未选择文件"

    filename = getattr(
        uploaded_file,
        "name",
        ""
    )

    extension = get_file_extension(
        filename
    )

    if not extension:
        return (
            False,
            "仅支持 TXT、MD、PDF、DOCX 文件"
        )

    try:
        size = uploaded_file.size
    except Exception:
        size = None

    if size == 0:
        return (
            False,
            "文件内容为空"
        )

    # 最大 20MB
    max_size = 20 * 1024 * 1024

    if (
        size is not None
        and size > max_size
    ):
        return (
            False,
            "文件超过 20MB"
        )

    return True, ""


def extract_requirement(uploaded_file):
    """
    根据文件格式自动解析需求。
    """

    valid, message = validate_uploaded_file(
        uploaded_file
    )

    if not valid:
        raise ValueError(message)

    filename = uploaded_file.name.lower()

    if filename.endswith(
        (".txt", ".md")
    ):
        text = read_text_file(
            uploaded_file
        )

    elif filename.endswith(".pdf"):
        text = read_pdf_file(
            uploaded_file
        )

    elif filename.endswith(".docx"):
        text = read_word_file(
            uploaded_file
        )

    else:
        raise ValueError(
            "暂不支持当前文件"
        )

    text = text.strip()

    if not text:
        raise ValueError(
            "文件中没有提取到有效文字"
        )

    return text