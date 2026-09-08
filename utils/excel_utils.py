# utils package
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import (
    Font,
    Alignment,
    Border,
    Side,
)


def create_border():
    """
    创建 Excel 边框。
    """

    thin = Side(
        style="thin"
    )

    return Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin,
    )


def style_header(
    worksheet,
    border
):
    """
    设置表头格式。
    """

    for cell in worksheet[1]:

        cell.font = Font(
            bold=True
        )

        cell.border = border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
            wrap_text=True
        )


def style_body(
    worksheet,
    border
):
    """
    设置内容格式。
    """

    for row in worksheet.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


def set_column_widths(
    worksheet,
    widths
):
    """
    设置 Excel 列宽。
    """

    for column, width in widths.items():

        worksheet.column_dimensions[
            column
        ].width = width


def create_test_result_excel(
    test_cases,
    rtm_result=None
):
    """
    生成测试结果 Excel。

    Sheet1：测试用例
    Sheet2：RTM需求追踪矩阵
    """

    if rtm_result is None:
        rtm_result = []

    workbook = Workbook()

    border = create_border()

    # ==============================================
    # Sheet1：测试用例
    # ==============================================

    worksheet = workbook.active

    worksheet.title = "测试用例"

    headers = [
        "用例编号",
        "测试模块",
        "测试标题",
        "前置条件",
        "操作步骤",
        "测试数据",
        "预期结果",
        "优先级",
    ]

    worksheet.append(headers)

    for case in test_cases:

        worksheet.append(
            [
                case.get(
                    "case_id",
                    ""
                ),

                case.get(
                    "module",
                    ""
                ),

                case.get(
                    "title",
                    ""
                ),

                case.get(
                    "precondition",
                    ""
                ),

                case.get(
                    "steps",
                    ""
                ),

                case.get(
                    "test_data",
                    ""
                ),

                case.get(
                    "expected_result",
                    ""
                ),

                case.get(
                    "priority",
                    ""
                ),
            ]
        )

    style_header(
        worksheet,
        border
    )

    style_body(
        worksheet,
        border
    )

    set_column_widths(
        worksheet,
        {
            "A": 12,
            "B": 18,
            "C": 35,
            "D": 30,
            "E": 50,
            "F": 35,
            "G": 45,
            "H": 10,
        }
    )

    worksheet.freeze_panes = "A2"

    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    # ==============================================
    # Sheet2：RTM
    # ==============================================

    rtm_sheet = workbook.create_sheet(
        "RTM需求追踪矩阵"
    )

    rtm_headers = [
        "需求编号",
        "需求内容",
        "需求类型",
        "关联用例",
        "覆盖状态",
        "备注",
    ]

    rtm_sheet.append(
        rtm_headers
    )

    for item in rtm_result:

        covered_cases = item.get(
            "covered_cases",
            []
        )

        if isinstance(
            covered_cases,
            list
        ):

            covered_cases = ", ".join(
                str(case_id)
                for case_id
                in covered_cases
            )

        rtm_sheet.append(
            [
                item.get(
                    "requirement_id",
                    ""
                ),

                item.get(
                    "requirement",
                    ""
                ),

                item.get(
                    "requirement_type",
                    ""
                ),

                covered_cases,

                item.get(
                    "coverage_status",
                    ""
                ),

                item.get(
                    "note",
                    ""
                ),
            ]
        )

    style_header(
        rtm_sheet,
        border
    )

    style_body(
        rtm_sheet,
        border
    )

    set_column_widths(
        rtm_sheet,
        {
            "A": 14,
            "B": 50,
            "C": 18,
            "D": 30,
            "E": 14,
            "F": 45,
        }
    )

    rtm_sheet.freeze_panes = "A2"

    rtm_sheet.auto_filter.ref = (
        rtm_sheet.dimensions
    )

    # ==============================================
    # 输出
    # ==============================================

    output = BytesIO()

    workbook.save(output)

    output.seek(0)

    return output.getvalue()