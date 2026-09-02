import os
import json
from io import BytesIO
from datetime import datetime

import pandas as pd
import streamlit as st

from dotenv import load_dotenv
from langchain_openai import ChatOpenAI

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side

from pypdf import PdfReader
from docx import Document

from rag import (
    get_rag_context,
    rebuild_vector_db,
    get_knowledge_file_list,
    get_knowledge_file_content,
    delete_knowledge_file
)


# =========================================================
# 1. 页面配置
# =========================================================

st.set_page_config(
    page_title="AI 软件测试平台",
    page_icon="🧪",
    layout="wide",
    initial_sidebar_state="expanded"
)


# =========================================================
# 2. 前端样式
# =========================================================

st.markdown(
    """
<style>

/* 页面整体 */
.stApp {
    background:
        linear-gradient(
            180deg,
            #f7f9fc 0%,
            #ffffff 320px
        );
}

.block-container {
    max-width: 1500px;
    padding-top: 1.3rem;
    padding-bottom: 4rem;
}

/* 隐藏 footer */
footer {
    visibility: hidden;
}

/* 顶部 Hero */
.hero-card {
    padding: 26px 30px;
    border-radius: 18px;
    border: 1px solid #e6eaf0;
    background:
        linear-gradient(
            135deg,
            #ffffff,
            #f4f7ff
        );
    box-shadow:
        0 6px 24px rgba(15, 23, 42, 0.05);
    margin-bottom: 20px;
}

.hero-title {
    font-size: 32px;
    font-weight: 800;
    letter-spacing: -0.6px;
    margin-bottom: 6px;
}

.hero-description {
    color: #667085;
    font-size: 15px;
    line-height: 1.7;
    margin-bottom: 14px;
}

.feature-chip {
    display: inline-block;
    padding: 5px 11px;
    border-radius: 20px;
    background: #eef4ff;
    color: #175cd3;
    font-size: 12px;
    font-weight: 650;
    margin-right: 6px;
    margin-top: 4px;
}

/* section */
.section-title {
    font-size: 22px;
    font-weight: 760;
    margin-top: 14px;
    margin-bottom: 6px;
}

.section-desc {
    color: #667085;
    font-size: 14px;
    margin-bottom: 14px;
}

/* Metric */
[data-testid="stMetric"] {
    background: white;
    border: 1px solid #e6eaf0;
    border-radius: 14px;
    padding: 14px 16px;
    box-shadow:
        0 3px 12px rgba(15, 23, 42, 0.035);
}

[data-testid="stMetricValue"] {
    font-size: 27px;
    font-weight: 760;
}

/* button */
.stButton > button {
    border-radius: 10px;
    min-height: 42px;
    font-weight: 650;
}

.stDownloadButton > button {
    border-radius: 10px;
    min-height: 42px;
    font-weight: 650;
}

/* input */
div[data-baseweb="input"] > div {
    border-radius: 10px;
}

textarea {
    border-radius: 10px !important;
}

/* alert */
[data-testid="stAlert"] {
    border-radius: 12px;
}

/* dataframe */
[data-testid="stDataFrame"] {
    border: 1px solid #e6eaf0;
    border-radius: 12px;
    overflow: hidden;
}

/* sidebar */
[data-testid="stSidebar"] {
    background: #f8fafc;
    border-right: 1px solid #e6eaf0;
}

/* tabs */
button[data-baseweb="tab"] {
    font-weight: 650;
}

</style>
""",
    unsafe_allow_html=True
)


# =========================================================
# 3. 路径
# =========================================================

BASE_DIR = os.path.dirname(
    os.path.abspath(__file__)
)

KNOWLEDGE_DIR = os.path.join(
    BASE_DIR,
    "knowledge"
)

HISTORY_BUG_FILE = os.path.join(
    KNOWLEDGE_DIR,
    "历史Bug.txt"
)

os.makedirs(
    KNOWLEDGE_DIR,
    exist_ok=True
)


# =========================================================
# 4. API 配置
# =========================================================

load_dotenv()

api_key = os.getenv(
    "DEEPSEEK_API_KEY"
)

if not api_key:

    st.error(
        "系统 API 配置缺失，请检查 .env 文件。"
    )

    st.stop()


model = ChatOpenAI(
    model="deepseek-chat",
    api_key=api_key,
    base_url="https://api.deepseek.com",
    temperature=0
)


# =========================================================
# 5. Session State
# =========================================================

default_states = {
    "test_cases": [],
    "review_content": "",
    "analysis_content": "",
    "rag_context": "",
    "current_requirement": "",
    "quality_result": None,
    "rtm_result": [],
    "requirement_items": [],
    "bug_analysis": None
}

for key, value in default_states.items():

    if key not in st.session_state:

        st.session_state[key] = value


# =========================================================
# 6. JSON 工具
# =========================================================

def clean_json_text(content):

    if not content:
        return ""

    content = content.strip()

    if content.startswith("```json"):
        content = content[len("```json"):].strip()

    elif content.startswith("```"):
        content = content[3:].strip()

    if content.endswith("```"):
        content = content[:-3].strip()

    return content


def parse_json_list(content):

    cleaned = clean_json_text(
        content
    )

    try:

        result = json.loads(
            cleaned
        )

        if isinstance(
            result,
            list
        ):
            return result

    except Exception:
        pass


    start = cleaned.find(
        "["
    )

    end = cleaned.rfind(
        "]"
    )

    if (
        start != -1
        and end != -1
        and end > start
    ):

        try:

            result = json.loads(
                cleaned[
                    start:end + 1
                ]
            )

            if isinstance(
                result,
                list
            ):
                return result

        except Exception:
            pass


    return None


def parse_json_dict(content):

    cleaned = clean_json_text(
        content
    )

    try:

        result = json.loads(
            cleaned
        )

        if isinstance(
            result,
            dict
        ):
            return result

    except Exception:
        pass


    start = cleaned.find(
        "{"
    )

    end = cleaned.rfind(
        "}"
    )

    if (
        start != -1
        and end != -1
        and end > start
    ):

        try:

            result = json.loads(
                cleaned[
                    start:end + 1
                ]
            )

            if isinstance(
                result,
                dict
            ):
                return result

        except Exception:
            pass


    return None


# =========================================================
# 7. 文件读取
# =========================================================

def read_text_file(uploaded_file):

    content = uploaded_file.getvalue()

    try:

        return content.decode(
            "utf-8"
        )

    except UnicodeDecodeError:

        return content.decode(
            "gbk",
            errors="ignore"
        )


def read_pdf_file(uploaded_file):

    text = ""

    try:

        reader = PdfReader(
            uploaded_file
        )

        for page_number, page in enumerate(
            reader.pages,
            start=1
        ):

            page_text = (
                page.extract_text()
            )

            if page_text:

                text += (
                    f"\n===== 第 {page_number} 页 =====\n"
                )

                text += page_text

    except Exception as e:

        st.error(
            f"PDF 读取失败：{e}"
        )

    return text


def read_word_file(uploaded_file):

    text = ""

    try:

        document = Document(
            uploaded_file
        )

        for paragraph in document.paragraphs:

            value = (
                paragraph.text.strip()
            )

            if value:

                text += (
                    value
                    + "\n"
                )


        for table in document.tables:

            text += (
                "\n===== 表格 =====\n"
            )

            for row in table.rows:

                text += (
                    " | ".join(
                        cell.text.strip()
                        for cell in row.cells
                    )
                    + "\n"
                )

    except Exception as e:

        st.error(
            f"Word 读取失败：{e}"
        )

    return text


def extract_requirement(uploaded_file):

    if uploaded_file is None:
        return ""

    name = (
        uploaded_file.name.lower()
    )

    if name.endswith(
        (".txt", ".md")
    ):

        return read_text_file(
            uploaded_file
        )

    if name.endswith(
        ".pdf"
    ):

        return read_pdf_file(
            uploaded_file
        )

    if name.endswith(
        ".docx"
    ):

        return read_word_file(
            uploaded_file
        )

    return ""


# =========================================================
# 8. 测试用例工具
# =========================================================

def normalize_case_ids(test_cases):

    for index, case in enumerate(
        test_cases,
        start=1
    ):

        case[
            "case_id"
        ] = f"TC{index:03d}"

    return test_cases


def test_cases_to_dataframe(test_cases):

    rows = []

    for case in test_cases:

        rows.append(
            {
                "用例编号":
                    case.get(
                        "case_id",
                        ""
                    ),

                "测试模块":
                    case.get(
                        "module",
                        ""
                    ),

                "测试标题":
                    case.get(
                        "title",
                        ""
                    ),

                "前置条件":
                    case.get(
                        "precondition",
                        ""
                    ),

                "操作步骤":
                    case.get(
                        "steps",
                        ""
                    ),

                "测试数据":
                    case.get(
                        "test_data",
                        ""
                    ),

                "预期结果":
                    case.get(
                        "expected_result",
                        ""
                    ),

                "优先级":
                    case.get(
                        "priority",
                        ""
                    )
            }
        )

    return pd.DataFrame(
        rows
    )


def dataframe_to_test_cases(dataframe):

    cases = []

    for _, row in dataframe.iterrows():

        cases.append(
            {
                "case_id":
                    str(
                        row.get(
                            "用例编号",
                            ""
                        )
                    ),

                "module":
                    str(
                        row.get(
                            "测试模块",
                            ""
                        )
                    ),

                "title":
                    str(
                        row.get(
                            "测试标题",
                            ""
                        )
                    ),

                "precondition":
                    str(
                        row.get(
                            "前置条件",
                            ""
                        )
                    ),

                "steps":
                    str(
                        row.get(
                            "操作步骤",
                            ""
                        )
                    ),

                "test_data":
                    str(
                        row.get(
                            "测试数据",
                            ""
                        )
                    ),

                "expected_result":
                    str(
                        row.get(
                            "预期结果",
                            ""
                        )
                    ),

                "priority":
                    str(
                        row.get(
                            "优先级",
                            ""
                        )
                    )
            }
        )

    return cases


# =========================================================
# 9. Excel
# =========================================================

def create_excel(
    test_cases,
    rtm_result
):

    wb = Workbook()

    thin = Side(
        style="thin"
    )

    border = Border(
        left=thin,
        right=thin,
        top=thin,
        bottom=thin
    )


    # =====================================================
    # Sheet 1 测试用例
    # =====================================================

    ws = wb.active

    ws.title = "测试用例"

    headers = [
        "用例编号",
        "测试模块",
        "测试标题",
        "前置条件",
        "操作步骤",
        "测试数据",
        "预期结果",
        "优先级"
    ]

    ws.append(
        headers
    )


    for case in test_cases:

        ws.append(
            [
                case.get(
                    "case_id",
                    ""
                ),
                case.get(
                    "module",
                    ""
                ),
                case.get(
                    "title",
                    ""
                ),
                case.get(
                    "precondition",
                    ""
                ),
                case.get(
                    "steps",
                    ""
                ),
                case.get(
                    "test_data",
                    ""
                ),
                case.get(
                    "expected_result",
                    ""
                ),
                case.get(
                    "priority",
                    ""
                )
            ]
        )


    for cell in ws[1]:

        cell.font = Font(
            bold=True
        )

        cell.border = border

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    for row in ws.iter_rows(
        min_row=2
    ):

        for cell in row:

            cell.border = border

            cell.alignment = Alignment(
                vertical="top",
                wrap_text=True
            )


    widths = {
        "A": 12,
        "B": 16,
        "C": 32,
        "D": 30,
        "E": 50,
        "F": 35,
        "G": 45,
        "H": 10
    }

    for column, width in (
        widths.items()
    ):

        ws.column_dimensions[
            column
        ].width = width


    ws.freeze_panes = "A2"

    ws.auto_filter.ref = (
        ws.dimensions
    )


    # =====================================================
    # Sheet 2 RTM
    # =====================================================

    if rtm_result:

        rtm_ws = wb.create_sheet(
            "需求追踪矩阵"
        )

        rtm_ws.append(
            [
                "需求编号",
                "需求内容",
                "需求类型",
                "覆盖用例",
                "覆盖状态",
                "说明"
            ]
        )


        for item in rtm_result:

            covered_cases = item.get(
                "covered_cases",
                []
            )

            if isinstance(
                covered_cases,
                list
            ):

                covered_cases = (
                    "、".join(
                        covered_cases
                    )
                )


            rtm_ws.append(
                [
                    item.get(
                        "requirement_id",
                        ""
                    ),
                    item.get(
                        "requirement",
                        ""
                    ),
                    item.get(
                        "requirement_type",
                        ""
                    ),
                    covered_cases,
                    item.get(
                        "coverage_status",
                        ""
                    ),
                    item.get(
                        "note",
                        ""
                    )
                ]
            )


        for cell in rtm_ws[1]:

            cell.font = Font(
                bold=True
            )

            cell.border = border

            cell.alignment = Alignment(
                horizontal="center",
                vertical="center"
            )


        for row in rtm_ws.iter_rows(
            min_row=2
        ):

            for cell in row:

                cell.border = border

                cell.alignment = Alignment(
                    vertical="top",
                    wrap_text=True
                )


        rtm_widths = {
            "A": 12,
            "B": 50,
            "C": 18,
            "D": 30,
            "E": 15,
            "F": 45
        }

        for column, width in (
            rtm_widths.items()
        ):

            rtm_ws.column_dimensions[
                column
            ].width = width


        rtm_ws.freeze_panes = "A2"


    output = BytesIO()

    wb.save(
        output
    )

    output.seek(
        0
    )

    return output


# =========================================================
# 10. Sidebar
# =========================================================

with st.sidebar:

    st.markdown(
        "## 🧪 AI 软件测试平台"
    )

    st.caption(
        "智能测试设计与质量分析"
    )

    st.divider()

    st.markdown(
        "### 📚 测试知识库"
    )

    knowledge_files = (
        st.file_uploader(
            "上传知识文件",
            type=[
                "txt",
                "md",
                "pdf",
                "docx"
            ],
            accept_multiple_files=True,
            key="knowledge_upload"
        )
    )


    if knowledge_files:

        if st.button(
            "保存并更新知识库",
            use_container_width=True
        ):

            try:

                for uploaded_file in (
                    knowledge_files
                ):

                    safe_name = (
                        os.path.basename(
                            uploaded_file.name
                        )
                    )

                    save_path = (
                        os.path.join(
                            KNOWLEDGE_DIR,
                            safe_name
                        )
                    )

                    with open(
                        save_path,
                        "wb"
                    ) as f:

                        f.write(
                            uploaded_file.getbuffer()
                        )


                with st.spinner(
                    "正在更新知识库..."
                ):

                    rebuild_vector_db()


                st.success(
                    "知识库更新成功"
                )

                st.rerun()

            except Exception as e:

                st.error(
                    f"知识库更新失败：{e}"
                )


    st.divider()


    try:

        knowledge_file_list = (
            get_knowledge_file_list()
        )

    except Exception:

        knowledge_file_list = []


    if knowledge_file_list:

        selected_file = (
            st.selectbox(
                "当前知识文件",
                knowledge_file_list
            )
        )


        preview_col1, preview_col2 = (
            st.columns(2)
        )


        with preview_col1:

            if st.button(
                "查看",
                use_container_width=True
            ):

                st.session_state[
                    "knowledge_preview_name"
                ] = selected_file

                st.session_state[
                    "knowledge_preview"
                ] = (
                    get_knowledge_file_content(
                        selected_file
                    )
                )


        with preview_col2:

            if st.button(
                "删除",
                use_container_width=True
            ):

                if delete_knowledge_file(
                    selected_file
                ):

                    with st.spinner(
                        "正在更新知识库..."
                    ):

                        rebuild_vector_db()

                    st.session_state.pop(
                        "knowledge_preview",
                        None
                    )

                    st.session_state.pop(
                        "knowledge_preview_name",
                        None
                    )

                    st.rerun()

    else:

        st.info(
            "知识库暂无文件"
        )


    st.divider()

    st.markdown(
        "### 📊 当前测试状态"
    )


    sidebar_cases = (
        st.session_state[
            "test_cases"
        ]
    )


    total_sidebar = len(
        sidebar_cases
    )


    p0_sidebar = sum(
        case.get(
            "priority"
        ) == "P0"
        for case in sidebar_cases
    )


    p1_sidebar = sum(
        case.get(
            "priority"
        ) == "P1"
        for case in sidebar_cases
    )


    p2_sidebar = sum(
        case.get(
            "priority"
        ) == "P2"
        for case in sidebar_cases
    )


    side1, side2 = st.columns(
        2
    )

    side1.metric(
        "用例",
        total_sidebar
    )

    side2.metric(
        "P0",
        p0_sidebar
    )


    side3, side4 = st.columns(
        2
    )

    side3.metric(
        "P1",
        p1_sidebar
    )

    side4.metric(
        "P2",
        p2_sidebar
    )


# =========================================================
# 11. Hero
# =========================================================

st.markdown(
    """
<div class="hero-card">

<div class="hero-title">
🧪 AI 软件测试平台
</div>

<div class="hero-description">
智能完成需求分析、测试点设计、测试用例生成、
质量评估、需求追踪、测试风险分析与经验沉淀。
</div>

<span class="feature-chip">需求评审</span>
<span class="feature-chip">测试设计</span>
<span class="feature-chip">质量评估</span>
<span class="feature-chip">需求追踪</span>
<span class="feature-chip">Bug 分析</span>

</div>
""",
    unsafe_allow_html=True
)


# =========================================================
# 12. 知识库预览
# =========================================================

if "knowledge_preview" in st.session_state:

    with st.expander(
        "📖 知识库文件预览"
    ):

        st.caption(
            st.session_state.get(
                "knowledge_preview_name",
                ""
            )
        )

        st.text_area(
            "文件内容",
            value=st.session_state[
                "knowledge_preview"
            ],
            height=300,
            disabled=True
        )


# =========================================================
# 13. 需求输入
# =========================================================

st.markdown(
    '<div class="section-title">'
    '① 需求输入'
    '</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="section-desc">'
    '输入产品需求，或上传需求文档进行智能测试分析。'
    '</div>',
    unsafe_allow_html=True
)


with st.container(
    border=True
):

    input_method = st.radio(
        "需求输入方式",
        [
            "手动输入",
            "上传需求文档"
        ],
        horizontal=True
    )


    requirement = ""


    if input_method == "手动输入":

        requirement = st.text_area(
            "需求内容",
            height=260,
            placeholder="""
例如：

用户通过手机号和密码登录。

手机号不能为空。
密码不能为空。

手机号必须为11位数字。

密码连续输错5次后，
账号锁定30分钟。

登录成功后进入首页。
"""
        )


    else:

        uploaded_requirement = (
            st.file_uploader(
                "上传需求文档",
                type=[
                    "txt",
                    "md",
                    "pdf",
                    "docx"
                ],
                key="requirement_upload"
            )
        )


        if uploaded_requirement:

            parsed_requirement = (
                extract_requirement(
                    uploaded_requirement
                )
            )

            requirement = st.text_area(
                "解析后的需求内容",
                value=parsed_requirement,
                height=280
            )


    start_col1, start_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with start_col1:

        start_analysis = st.button(
            "🚀 开始智能测试分析",
            type="primary",
            use_container_width=True
        )


    with start_col2:

        st.caption(
            "系统将依次完成：需求评审 → 测试点 → 测试用例"
        )


# =========================================================
# 14. 主分析流程
# =========================================================

if start_analysis:

    if not requirement.strip():

        st.warning(
            "请先输入需求内容。"
        )

        st.stop()


    st.session_state[
        "current_requirement"
    ] = requirement

    st.session_state[
        "quality_result"
    ] = None

    st.session_state[
        "rtm_result"
    ] = []

    st.session_state[
        "requirement_items"
    ] = []


    # =====================================================
    # RAG
    # =====================================================

    try:

        with st.spinner(
            "正在检索测试经验..."
        ):

            rag_context = (
                get_rag_context(
                    requirement,
                    k=4
                )
            )

    except Exception as e:

        rag_context = (
            f"知识检索失败：{e}"
        )


    st.session_state[
        "rag_context"
    ] = rag_context


    # =====================================================
    # 需求评审
    # =====================================================

    review_prompt = f"""
你是一名资深软件测试工程师。

请严格根据当前需求进行需求评审。

规则：

1. 当前需求是唯一业务依据。
2. 历史测试资料只作为经验参考。
3. 不允许自行创造需求规则。
4. 未说明内容必须标记【需求待确认】。

输出：

## 一、需求理解
## 二、已明确需求规则
## 三、需求风险
## 四、需求待确认
## 五、测试建议

当前需求：

{requirement}

历史测试参考：

{rag_context}
"""


    try:

        with st.spinner(
            "正在进行需求评审..."
        ):

            review_response = (
                model.invoke(
                    review_prompt
                )
            )

    except Exception as e:

        st.error(
            f"需求评审失败：{e}"
        )

        st.stop()


    review_content = (
        review_response.content
    )

    st.session_state[
        "review_content"
    ] = review_content


    # =====================================================
    # 测试点
    # =====================================================

    analysis_prompt = f"""
你是一名专业软件测试工程师。

根据需求和需求评审分析完整测试点。

必须覆盖：

1. 正常场景
2. 异常场景
3. 边界值
4. 业务规则
5. 数据校验
6. 状态流转
7. 权限与安全
8. 兼容性
9. 历史缺陷回归
10. 需求待确认

要求：

不得创造需求中没有的业务规则。

需求：

{requirement}

需求评审：

{review_content}

测试经验：

{rag_context}
"""


    try:

        with st.spinner(
            "正在设计测试点..."
        ):

            analysis_response = (
                model.invoke(
                    analysis_prompt
                )
            )

    except Exception as e:

        st.error(
            f"测试点分析失败：{e}"
        )

        st.stop()


    analysis_content = (
        analysis_response.content
    )

    st.session_state[
        "analysis_content"
    ] = analysis_content


    # =====================================================
    # 测试用例
    # =====================================================

    case_prompt = f"""
你是一名专业软件测试工程师。

根据需求、需求评审、测试点和历史测试经验
生成完整测试用例。

字段必须包含：

case_id
module
title
precondition
steps
test_data
expected_result
priority

priority 只能：

P0
P1
P2

要求：

1. 覆盖正常场景。
2. 覆盖异常场景。
3. 覆盖边界场景。
4. 覆盖核心业务规则。
5. 历史 Bug 只作为回归经验。
6. 不允许创造需求中不存在的规则。
7. 未明确内容标记【需求待确认】。
8. 避免重复测试用例。
9. 只输出合法 JSON 数组。
10. 不输出 Markdown。

需求：

{requirement}

需求评审：

{review_content}

测试点：

{analysis_content}

历史测试经验：

{rag_context}
"""


    try:

        with st.spinner(
            "正在生成测试用例..."
        ):

            case_response = (
                model.invoke(
                    case_prompt
                )
            )

    except Exception as e:

        st.error(
            f"测试用例生成失败：{e}"
        )

        st.stop()


    test_cases = parse_json_list(
        case_response.content
    )


    if test_cases is None:

        st.error(
            "测试用例数据解析失败。"
        )

        with st.expander(
            "查看原始返回内容"
        ):

            st.code(
                case_response.content
            )

        st.stop()


    st.session_state[
        "test_cases"
    ] = normalize_case_ids(
        test_cases
    )


    st.success(
        f"分析完成，共生成 {len(test_cases)} 条测试用例。"
    )


# =========================================================
# 15. 分析结果
# =========================================================

if (
    st.session_state[
        "review_content"
    ]
    or st.session_state[
        "analysis_content"
    ]
):

    st.divider()

    st.markdown(
        '<div class="section-title">'
        '② 智能分析结果'
        '</div>',
        unsafe_allow_html=True
    )


    review_tab, point_tab, knowledge_tab = (
        st.tabs(
            [
                "📋 需求评审",
                "🎯 测试点",
                "📚 测试经验参考"
            ]
        )
    )


    with review_tab:

        st.markdown(
            st.session_state[
                "review_content"
            ]
        )


    with point_tab:

        st.markdown(
            st.session_state[
                "analysis_content"
            ]
        )


    with knowledge_tab:

        st.text(
            st.session_state[
                "rag_context"
            ]
        )


# =========================================================
# 16. 测试用例
# =========================================================

if st.session_state[
    "test_cases"
]:

    st.divider()

    st.markdown(
        '<div class="section-title">'
        '③ 测试用例'
        '</div>',
        unsafe_allow_html=True
    )

    st.caption(
        "支持直接编辑、新增和删除测试用例。"
    )


    dataframe = (
        test_cases_to_dataframe(
            st.session_state[
                "test_cases"
            ]
        )
    )


    edited_df = st.data_editor(
        dataframe,
        use_container_width=True,
        hide_index=True,
        num_rows="dynamic",
        height=450,
        key="test_case_editor",
        column_config={
            "用例编号":
                st.column_config.TextColumn(
                    "用例编号",
                    width="small"
                ),

            "测试标题":
                st.column_config.TextColumn(
                    "测试标题",
                    width="large"
                ),

            "优先级":
                st.column_config.SelectboxColumn(
                    "优先级",
                    options=[
                        "P0",
                        "P1",
                        "P2"
                    ],
                    width="small"
                )
        }
    )


    current_cases = (
        dataframe_to_test_cases(
            edited_df
        )
    )


    save_col1, save_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with save_col1:

        save_cases = st.button(
            "💾 保存修改",
            use_container_width=True
        )


    if save_cases:

        current_cases = (
            normalize_case_ids(
                current_cases
            )
        )

        st.session_state[
            "test_cases"
        ] = current_cases

        st.session_state[
            "quality_result"
        ] = None

        st.session_state[
            "rtm_result"
        ] = []

        st.success(
            "测试用例修改已保存。"
        )

        st.rerun()


    current_requirement = (
        st.session_state[
            "current_requirement"
        ]
    )

    current_analysis = (
        st.session_state[
            "analysis_content"
        ]
    )

    current_rag = (
        st.session_state[
            "rag_context"
        ]
    )


    # =====================================================
    # 17. 智能优化
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '④ 智能用例优化'
        '</div>',
        unsafe_allow_html=True
    )

    st.markdown(
        '<div class="section-desc">'
        '检查遗漏、重复及用例设计质量。'
        '</div>',
        unsafe_allow_html=True
    )


    optimize_col1, optimize_col2, optimize_col3 = (
        st.columns(3)
    )


    with optimize_col1:

        add_missing = st.button(
            "➕ 补充遗漏用例",
            use_container_width=True
        )


    with optimize_col2:

        remove_duplicate = st.button(
            "🧹 删除重复用例",
            use_container_width=True
        )


    with optimize_col3:

        optimize_all = st.button(
            "✨ 全面优化用例",
            type="primary",
            use_container_width=True
        )


    if add_missing:

        prompt = f"""
你是一名高级软件测试工程师。

检查现有测试用例是否存在遗漏。

要求：

1. 保留已有有效用例。
2. 只补充真正遗漏的场景。
3. 不生成重复用例。
4. 不创造需求规则。
5. 返回原用例 + 新增用例的完整 JSON 数组。

字段：

case_id
module
title
precondition
steps
test_data
expected_result
priority

需求：

{current_requirement}

测试点：

{current_analysis}

历史测试经验：

{current_rag}

现有用例：

{json.dumps(
    current_cases,
    ensure_ascii=False
)}
"""


        try:

            with st.spinner(
                "正在检查遗漏..."
            ):

                response = (
                    model.invoke(
                        prompt
                    )
                )

            optimized = (
                parse_json_list(
                    response.content
                )
            )

        except Exception as e:

            optimized = None

            st.error(
                f"用例补充失败：{e}"
            )


        if optimized:

            st.session_state[
                "test_cases"
            ] = normalize_case_ids(
                optimized
            )

            st.session_state[
                "quality_result"
            ] = None

            st.session_state[
                "rtm_result"
            ] = []

            st.rerun()


    if remove_duplicate:

        prompt = f"""
你是一名高级软件测试工程师。

删除下面测试用例中的重复或高度重复场景。

要求：

1. 不删除具有独立测试价值的用例。
2. 保留边界和异常测试。
3. 不创造新的需求规则。
4. 返回完整 JSON 数组。

测试用例：

{json.dumps(
    current_cases,
    ensure_ascii=False
)}
"""


        try:

            with st.spinner(
                "正在检查重复用例..."
            ):

                response = (
                    model.invoke(
                        prompt
                    )
                )

            optimized = (
                parse_json_list(
                    response.content
                )
            )

        except Exception as e:

            optimized = None

            st.error(
                f"用例去重失败：{e}"
            )


        if optimized:

            st.session_state[
                "test_cases"
            ] = normalize_case_ids(
                optimized
            )

            st.session_state[
                "quality_result"
            ] = None

            st.session_state[
                "rtm_result"
            ] = []

            st.rerun()


    if optimize_all:

        prompt = f"""
你是一名高级软件测试工程师。

全面优化当前测试用例。

目标：

1. 删除重复用例。
2. 补充遗漏测试场景。
3. 优化标题。
4. 优化前置条件。
5. 优化步骤。
6. 优化测试数据。
7. 优化预期结果。
8. 调整不合理优先级。
9. 不创造需求中不存在的规则。

字段：

case_id
module
title
precondition
steps
test_data
expected_result
priority

只输出 JSON 数组。

需求：

{current_requirement}

测试点：

{current_analysis}

历史经验：

{current_rag}

当前用例：

{json.dumps(
    current_cases,
    ensure_ascii=False
)}
"""


        try:

            with st.spinner(
                "正在全面优化测试用例..."
            ):

                response = (
                    model.invoke(
                        prompt
                    )
                )

            optimized = (
                parse_json_list(
                    response.content
                )
            )

        except Exception as e:

            optimized = None

            st.error(
                f"全面优化失败：{e}"
            )


        if optimized:

            st.session_state[
                "test_cases"
            ] = normalize_case_ids(
                optimized
            )

            st.session_state[
                "quality_result"
            ] = None

            st.session_state[
                "rtm_result"
            ] = []

            st.rerun()


    # =====================================================
    # 18. 质量评估
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '⑤ 测试质量评估'
        '</div>',
        unsafe_allow_html=True
    )

    st.caption(
        "从需求覆盖、边界、异常、业务规则和可执行性等维度进行评估。"
    )


    quality_col1, quality_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with quality_col1:

        quality_button = st.button(
            "📊 开始质量评估",
            type="primary",
            use_container_width=True
        )


    if quality_button:

        prompt = f"""
你是一名高级软件测试负责人。

严格评估当前测试用例质量。

返回合法 JSON 对象：

{{
  "overall_score": 0,
  "requirement_coverage": 0,
  "normal_coverage": 0,
  "exception_coverage": 0,
  "boundary_coverage": 0,
  "business_rule_coverage": 0,
  "executability": 0,
  "non_duplicate_score": 0,
  "strengths": [],
  "missing_scenarios": [],
  "duplicate_risks": [],
  "high_risk_gaps": [],
  "improvement_suggestions": [],
  "summary": ""
}}

所有评分范围 0-100。

不得创造需求中不存在的规则。

需求：

{current_requirement}

测试点：

{current_analysis}

测试用例：

{json.dumps(
    current_cases,
    ensure_ascii=False
)}
"""


        try:

            with st.spinner(
                "正在评估测试质量..."
            ):

                response = (
                    model.invoke(
                        prompt
                    )
                )

            quality_result = (
                parse_json_dict(
                    response.content
                )
            )

        except Exception as e:

            quality_result = None

            st.error(
                f"质量评估失败：{e}"
            )


        if quality_result:

            st.session_state[
                "quality_result"
            ] = quality_result


    quality_result = (
        st.session_state[
            "quality_result"
        ]
    )


    if quality_result:

        overall_score = int(
            quality_result.get(
                "overall_score",
                0
            )
        )


        quality1, quality2 = (
            st.columns(
                [
                    1,
                    3
                ]
            )
        )


        with quality1:

            st.metric(
                "总体质量分",
                f"{overall_score}/100"
            )


        with quality2:

            st.progress(
                min(
                    max(
                        overall_score,
                        0
                    ),
                    100
                )
            )

            st.write(
                quality_result.get(
                    "summary",
                    ""
                )
            )


        metric_items = [
            (
                "需求覆盖",
                "requirement_coverage"
            ),
            (
                "正常场景",
                "normal_coverage"
            ),
            (
                "异常场景",
                "exception_coverage"
            ),
            (
                "边界覆盖",
                "boundary_coverage"
            ),
            (
                "业务规则",
                "business_rule_coverage"
            ),
            (
                "可执行性",
                "executability"
            ),
            (
                "非重复",
                "non_duplicate_score"
            )
        ]


        quality_columns = st.columns(
            7
        )


        for column, (
            label,
            key
        ) in zip(
            quality_columns,
            metric_items
        ):

            column.metric(
                label,
                quality_result.get(
                    key,
                    0
                )
            )


        quality_tab1, quality_tab2, quality_tab3, quality_tab4 = (
            st.tabs(
                [
                    "✅ 优势",
                    "⚠️ 遗漏",
                    "🚨 高风险",
                    "💡 优化建议"
                ]
            )
        )


        with quality_tab1:

            for item in quality_result.get(
                "strengths",
                []
            ):

                st.write(
                    f"• {item}"
                )


        with quality_tab2:

            items = quality_result.get(
                "missing_scenarios",
                []
            )

            if items:

                for item in items:

                    st.write(
                        f"• {item}"
                    )

            else:

                st.success(
                    "未发现明显遗漏"
                )


        with quality_tab3:

            items = quality_result.get(
                "high_risk_gaps",
                []
            )

            if items:

                for item in items:

                    st.write(
                        f"• {item}"
                    )

            else:

                st.success(
                    "未发现明显高风险遗漏"
                )


        with quality_tab4:

            for item in quality_result.get(
                "improvement_suggestions",
                []
            ):

                st.write(
                    f"• {item}"
                )


    # =====================================================
    # 19. RTM
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '⑥ 需求追踪矩阵'
        '</div>',
        unsafe_allow_html=True
    )

    st.caption(
        "检查每条需求是否被测试用例真正覆盖。"
    )


    rtm_col1, rtm_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with rtm_col1:

        generate_rtm = st.button(
            "🔗 生成需求追踪矩阵",
            use_container_width=True
        )


    if generate_rtm:

        requirement_prompt = f"""
你是一名软件测试需求分析工程师。

把原始需求拆分成最小可验证需求条目。

字段：

requirement_id
requirement
requirement_type

编号从 R001 开始。

不得创造需求中不存在的规则。

只输出 JSON 数组。

需求：

{current_requirement}
"""


        try:

            with st.spinner(
                "正在拆分需求..."
            ):

                response = (
                    model.invoke(
                        requirement_prompt
                    )
                )

            requirement_items = (
                parse_json_list(
                    response.content
                )
            )

        except Exception as e:

            requirement_items = None

            st.error(
                f"需求拆分失败：{e}"
            )


        if requirement_items:

            st.session_state[
                "requirement_items"
            ] = requirement_items


            rtm_prompt = f"""
你是一名高级软件测试工程师。

建立需求追踪矩阵。

字段：

requirement_id
requirement
requirement_type
covered_cases
coverage_status
note

coverage_status 只能：

已覆盖
部分覆盖
未覆盖

规则：

1. 不强行建立需求与用例映射。
2. covered_cases 只能使用已有测试用例编号。
3. 没有覆盖返回空数组。
4. 每条需求必须返回。
5. 不允许新增需求。

需求：

{json.dumps(
    requirement_items,
    ensure_ascii=False
)}

测试用例：

{json.dumps(
    current_cases,
    ensure_ascii=False
)}
"""


            try:

                with st.spinner(
                    "正在建立需求追踪关系..."
                ):

                    response = (
                        model.invoke(
                            rtm_prompt
                        )
                    )

                rtm_result = (
                    parse_json_list(
                        response.content
                    )
                )

            except Exception as e:

                rtm_result = None

                st.error(
                    f"需求追踪生成失败：{e}"
                )


            if rtm_result:

                st.session_state[
                    "rtm_result"
                ] = rtm_result

                st.rerun()


    rtm_result = (
        st.session_state[
            "rtm_result"
        ]
    )


    if rtm_result:

        rtm_rows = []


        for item in rtm_result:

            covered_cases = item.get(
                "covered_cases",
                []
            )

            if isinstance(
                covered_cases,
                list
            ):

                covered_cases = (
                    "、".join(
                        covered_cases
                    )
                )


            rtm_rows.append(
                {
                    "需求编号":
                        item.get(
                            "requirement_id",
                            ""
                        ),

                    "需求内容":
                        item.get(
                            "requirement",
                            ""
                        ),

                    "需求类型":
                        item.get(
                            "requirement_type",
                            ""
                        ),

                    "覆盖用例":
                        covered_cases,

                    "覆盖状态":
                        item.get(
                            "coverage_status",
                            ""
                        ),

                    "说明":
                        item.get(
                            "note",
                            ""
                        )
                }
            )


        st.dataframe(
            pd.DataFrame(
                rtm_rows
            ),
            use_container_width=True,
            hide_index=True
        )


    # =====================================================
    # 20. Dashboard
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '⑦ 测试质量 Dashboard'
        '</div>',
        unsafe_allow_html=True
    )


    total_count = len(
        current_cases
    )


    p0_count = sum(
        case.get(
            "priority"
        ) == "P0"
        for case in current_cases
    )


    p1_count = sum(
        case.get(
            "priority"
        ) == "P1"
        for case in current_cases
    )


    p2_count = sum(
        case.get(
            "priority"
        ) == "P2"
        for case in current_cases
    )


    total_requirements = len(
        rtm_result
    )


    covered_count = sum(
        item.get(
            "coverage_status"
        ) == "已覆盖"
        for item in rtm_result
    )


    partial_count = sum(
        item.get(
            "coverage_status"
        ) == "部分覆盖"
        for item in rtm_result
    )


    uncovered_count = sum(
        item.get(
            "coverage_status"
        ) == "未覆盖"
        for item in rtm_result
    )


    if total_requirements:

        coverage_rate = round(
            covered_count
            / total_requirements
            * 100,
            1
        )

    else:

        coverage_rate = 0


    if quality_result:

        overall_score = int(
            quality_result.get(
                "overall_score",
                0
            )
        )

        high_risk_gaps = (
            quality_result.get(
                "high_risk_gaps",
                []
            )
        )

        duplicate_risks = (
            quality_result.get(
                "duplicate_risks",
                []
            )
        )

    else:

        overall_score = 0

        high_risk_gaps = []

        duplicate_risks = []


    dashboard1, dashboard2, dashboard3, dashboard4 = (
        st.columns(4)
    )


    dashboard1.metric(
        "测试用例总数",
        total_count
    )

    dashboard2.metric(
        "总体质量分",
        f"{overall_score}/100"
    )

    dashboard3.metric(
        "需求完全覆盖率",
        f"{coverage_rate}%"
    )

    dashboard4.metric(
        "未覆盖需求",
        uncovered_count
    )


    dashboard5, dashboard6, dashboard7, dashboard8 = (
        st.columns(4)
    )


    dashboard5.metric(
        "P0 核心用例",
        p0_count
    )

    dashboard6.metric(
        "部分覆盖需求",
        partial_count
    )

    dashboard7.metric(
        "高风险遗漏",
        len(
            high_risk_gaps
        )
    )

    dashboard8.metric(
        "重复风险",
        len(
            duplicate_risks
        )
    )


    if (
        overall_score >= 90
        and coverage_rate >= 95
        and uncovered_count == 0
    ):

        st.success(
            "✅ 当前测试设计质量较高，可进入最终人工评审。"
        )

    elif (
        overall_score >= 75
        and coverage_rate >= 80
    ):

        st.warning(
            "⚠️ 当前测试设计基本可用，建议继续处理遗漏和高风险场景。"
        )

    else:

        st.info(
            "完成质量评估和需求追踪后，可获得更完整的测试质量判断。"
        )


    # =====================================================
    # 21. 导出
    # =====================================================

    st.markdown(
        '<div class="section-title">'
        '⑧ 导出测试结果'
        '</div>',
        unsafe_allow_html=True
    )


    excel_file = create_excel(
        normalize_case_ids(
            current_cases
        ),
        rtm_result
    )


    export_col1, export_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with export_col1:

        st.download_button(
            "⬇️ 下载测试结果",
            data=excel_file,
            file_name="AI测试结果.xlsx",
            mime=(
                "application/"
                "vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            ),
            type="primary",
            use_container_width=True
        )


# =========================================================
# 22. Bug 分析
# =========================================================

st.divider()

st.markdown(
    '<div class="section-title">'
    '⑨ Bug 智能分析'
    '</div>',
    unsafe_allow_html=True
)

st.markdown(
    '<div class="section-desc">'
    '分析实际测试缺陷，并将有效经验沉淀到测试知识库。'
    '</div>',
    unsafe_allow_html=True
)


with st.container(
    border=True
):

    bug_col1, bug_col2 = (
        st.columns(2)
    )


    with bug_col1:

        bug_title = st.text_input(
            "Bug 标题",
            placeholder="例如：连续输错5次密码后账号未锁定"
        )


        bug_description = st.text_area(
            "Bug 描述",
            height=120
        )


        bug_environment = st.text_input(
            "测试环境",
            placeholder="例如：Android 14 / 测试环境"
        )


    with bug_col2:

        bug_steps = st.text_area(
            "复现步骤",
            height=120
        )


        bug_expected = st.text_area(
            "预期结果",
            height=90
        )


        bug_actual = st.text_area(
            "实际结果",
            height=90
        )


    bug_btn_col1, bug_btn_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with bug_btn_col1:

        analyze_bug = st.button(
            "🔍 开始 Bug 分析",
            type="primary",
            use_container_width=True
        )


if analyze_bug:

    if not (
        bug_title.strip()
        or bug_description.strip()
    ):

        st.warning(
            "请至少填写 Bug 标题或 Bug 描述。"
        )

    else:

        bug_prompt = f"""
你是一名高级软件测试工程师。

分析下面的 Bug。

注意：

1. 不要声称已经确认真实代码根因。
2. root_causes 只能表示可能原因。
3. 给出实际可执行的定位建议。
4. 分析测试遗漏。
5. 给出回归测试点。
6. knowledge_summary 需要适合保存到历史 Bug 知识库。

只输出合法 JSON：

{{
  "severity": "S1/S2/S3/S4",
  "priority": "P0/P1/P2",
  "bug_type": "功能/接口/数据/兼容性/性能/安全/UI/其他",
  "module": "",
  "impact": "",
  "root_causes": [],
  "localization_steps": [],
  "test_gaps": [],
  "regression_points": [],
  "knowledge_summary": ""
}}

Bug标题：

{bug_title}

Bug描述：

{bug_description}

复现步骤：

{bug_steps}

预期结果：

{bug_expected}

实际结果：

{bug_actual}

环境：

{bug_environment}
"""


        try:

            with st.spinner(
                "正在分析 Bug..."
            ):

                response = (
                    model.invoke(
                        bug_prompt
                    )
                )

            bug_analysis = (
                parse_json_dict(
                    response.content
                )
            )

        except Exception as e:

            bug_analysis = None

            st.error(
                f"Bug 分析失败：{e}"
            )


        if bug_analysis:

            st.session_state[
                "bug_analysis"
            ] = bug_analysis


bug_analysis = (
    st.session_state[
        "bug_analysis"
    ]
)


if bug_analysis:

    st.markdown(
        "### Bug 分析结果"
    )


    bug_metric1, bug_metric2, bug_metric3 = (
        st.columns(3)
    )


    bug_metric1.metric(
        "严重程度",
        bug_analysis.get(
            "severity",
            "-"
        )
    )

    bug_metric2.metric(
        "优先级",
        bug_analysis.get(
            "priority",
            "-"
        )
    )

    bug_metric3.metric(
        "Bug 类型",
        bug_analysis.get(
            "bug_type",
            "-"
        )
    )


    bug_tab1, bug_tab2, bug_tab3, bug_tab4 = (
        st.tabs(
            [
                "🔍 可能原因",
                "🛠 定位建议",
                "⚠️ 测试遗漏",
                "🔁 回归测试"
            ]
        )
    )


    with bug_tab1:

        for item in bug_analysis.get(
            "root_causes",
            []
        ):

            st.write(
                f"• {item}"
            )


    with bug_tab2:

        for index, item in enumerate(
            bug_analysis.get(
                "localization_steps",
                []
            ),
            start=1
        ):

            st.write(
                f"{index}. {item}"
            )


    with bug_tab3:

        for item in bug_analysis.get(
            "test_gaps",
            []
        ):

            st.write(
                f"• {item}"
            )


    with bug_tab4:

        for item in bug_analysis.get(
            "regression_points",
            []
        ):

            st.write(
                f"• {item}"
            )


    knowledge_summary = st.text_area(
        "知识库摘要",
        value=bug_analysis.get(
            "knowledge_summary",
            ""
        ),
        height=150,
        key="knowledge_summary_editor"
    )


    if st.button(
        "📚 保存到历史 Bug 知识库"
    ):

        try:

            timestamp = (
                datetime.now().strftime(
                    "%Y-%m-%d %H:%M:%S"
                )
            )


            record = f"""

==================================================
历史 Bug
==================================================

记录时间：
{timestamp}

Bug标题：
{bug_title}

Bug描述：
{bug_description}

复现步骤：
{bug_steps}

预期结果：
{bug_expected}

实际结果：
{bug_actual}

测试环境：
{bug_environment}

严重程度：
{bug_analysis.get("severity", "")}

优先级：
{bug_analysis.get("priority", "")}

Bug类型：
{bug_analysis.get("bug_type", "")}

可能原因：
{chr(10).join("- " + str(item) for item in bug_analysis.get("root_causes", []))}

测试遗漏：
{chr(10).join("- " + str(item) for item in bug_analysis.get("test_gaps", []))}

回归测试点：
{chr(10).join("- " + str(item) for item in bug_analysis.get("regression_points", []))}

知识摘要：
{knowledge_summary}

==================================================

"""


            with open(
                HISTORY_BUG_FILE,
                "a",
                encoding="utf-8"
            ) as f:

                f.write(
                    record
                )


            with st.spinner(
                "正在更新测试知识库..."
            ):

                rebuild_vector_db()


            st.success(
                "Bug 已保存到历史测试知识库。"
            )


        except Exception as e:

            st.error(
                f"保存失败：{e}"
            )


# =========================================================
# 23. 清空
# =========================================================

if st.session_state[
    "test_cases"
]:

    st.divider()


    clear_col1, clear_col2 = (
        st.columns(
            [
                1,
                4
            ]
        )
    )


    with clear_col1:

        if st.button(
            "🗑 清空当前分析",
            use_container_width=True
        ):

            st.session_state[
                "test_cases"
            ] = []

            st.session_state[
                "review_content"
            ] = ""

            st.session_state[
                "analysis_content"
            ] = ""

            st.session_state[
                "rag_context"
            ] = ""

            st.session_state[
                "current_requirement"
            ] = ""

            st.session_state[
                "quality_result"
            ] = None

            st.session_state[
                "rtm_result"
            ] = []

            st.session_state[
                "requirement_items"
            ] = []

            st.rerun()